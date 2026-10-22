"""Dump an .xlsx file as stable, line-oriented text so git can diff it.

Usage: python xlsx2txt.py <file.xlsx>
Also works as a git textconv driver (git passes the file path as argv[1]).
Formulas are shown as written; cached values are shown after '=>' when available.
"""
import sys

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def dump(path):
    formulas = load_workbook(path, data_only=False)
    values = load_workbook(path, data_only=True)
    out = []
    for ws in formulas.worksheets:
        vs = values[ws.title]
        out.append(f"### Sheet: {ws.title} (dimensions {ws.dimensions})")
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                ref = f"{get_column_letter(cell.column)}{cell.row}"
                text = repr(cell.value) if not isinstance(cell.value, str) else cell.value
                cached = vs[cell.coordinate].value
                if isinstance(cell.value, str) and cell.value.startswith("=") and cached is not None:
                    text = f"{text}  => {cached!r}"
                out.append(f"{ws.title}!{ref}: {text}")
    return "\n".join(out) + "\n"


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8", newline="\n")
    sys.stdout.write(dump(sys.argv[1]))
